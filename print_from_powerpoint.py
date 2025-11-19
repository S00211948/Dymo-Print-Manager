import pandas as pd
from pptx import Presentation
import win32com.client
import os
from json import load as jsload
import sys

class PowerpointPrintService:
    # Read file paths
    config = None
    template_file = None
    output_folder = "C:/temp_invites"
    ppt = None
    
    def __init__(self):
        try:
            # Read config
            self.config = self.readConfig()
            self.template_file = self.config['pptx-template']
            os.makedirs(self.output_folder, exist_ok=True)
            # Create powerpoint client
            self.ppt = win32com.client.Dispatch("PowerPoint.Application")
            self.ppt.Visible = True  # optional, can show PowerPoint while printing

            print("> Powerpoint Print Manager Initialized")
        except Exception as e:
            print(f"> Powerpoint Print Manager Failed to Initialize\nError: {e}")

    # Functions
    def get_path(self):
        if hasattr(sys, 'frozen'):  
            # When running as .exe
            return os.path.dirname(os.path.realpath(sys.executable))
        else:
            # When running as .py
            return os.path.dirname(os.path.realpath(__file__))

    def readConfig(self):
        # Get the directory of the current file
        filepath = self.get_path()
        config_path = os.path.join(filepath,'config.json')
        with open(config_path,'r') as config:
            data = jsload(config)
        return data

    def printSlide(self, dataObj):
        # Concat names
        guest_names=''
        for i in range(1,7):
            if isinstance(dataObj[f'Guest_{i}'],str) and dataObj[f'Guest_{i}'] != '':
                if isinstance(dataObj[f'Guest_{i+1}'],str) and dataObj[f'Guest_{i+1}'] != '':
                    if i > 1:
                        guest_names +=', '
                else:
                    if i < 6:
                        guest_names += ', and '

                guest_names += dataObj[f'Guest_{i}']

        # Create invite from template
        prs = Presentation(self.template_file)
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    if "{{NAMES}}" in shape.text:
                        shape.text = shape.text.replace("{{NAMES}}", guest_names)

        temp_path = os.path.join(self.output_folder, f"Invite_{str.strip(str(dataObj['Employee']).replace(' ','_'))}_{str.strip(dataObj['ID'])}.pptx")
        print(temp_path)
        prs.save(temp_path)

        # Print via PowerPoint COM interface
        print(temp_path)
        presentation = self.ppt.Presentations.Open(temp_path, WithWindow=False)
        presentation.PrintOut()  # Default printer
        presentation.Close()